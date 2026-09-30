"""
JEVBP / JBESS Part Code Generator — v3
========================================
Flask + SQLite. Session-based auth, role/module access control, an
Admin-only category builder (add new material types at runtime), and a
JSON API consumed by a single-page app shell (templates/app.html).

Run:
    pip install -r requirements.txt
    python app.py
Then open http://localhost:5000  (default admin: Master / Master@123)
"""

import csv
import io
import json
import re
import sqlite3
import zipfile
from datetime import datetime
from functools import wraps
from xml.sax.saxutils import escape as xml_escape

from flask import Flask, render_template, request, redirect, url_for, session, jsonify, send_file
from werkzeug.security import generate_password_hash, check_password_hash

from rules import CATEGORIES as BUILTIN_CATEGORIES

DB_PATH = "partcodes.db"
# Modules an Admin can delegate to a "User" role account. "Manage Categories"
# is deliberately NOT in this list — creating new material types is kept
# Admin-only and isn't something Admins can hand off via access checkboxes.
ALL_MODULES = ["home", "history", "user_management"]

# Every generated code (hyphens included) must be this many characters long.
CODE_MIN_LEN = 10
CODE_MAX_LEN = 18

PLANT_INFO = {
    "name": "BESS Pune",
    "code": "4971",
    "address": "Kamshet, Pune, Maharashtra, India",
}
COMPANY_NAME = "JSW Energy PSP Eleven LTD."

app = Flask(__name__)
app.secret_key = "dev-secret-change-me"  # replace before real deployment


# --------------------------------------------------------------------------
# DB helpers
# --------------------------------------------------------------------------
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    conn.execute(
        """CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'User',
            access TEXT NOT NULL DEFAULT '[]',
            created_at TEXT NOT NULL
        )"""
    )
    conn.execute(
        """CREATE TABLE IF NOT EXISTS part_codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category TEXT NOT NULL,
            code TEXT NOT NULL UNIQUE,
            attributes TEXT NOT NULL,
            created_by TEXT,
            created_at TEXT NOT NULL
        )"""
    )
    conn.execute(
        """CREATE TABLE IF NOT EXISTS custom_categories (
            key TEXT PRIMARY KEY,
            label TEXT NOT NULL,
            code_format TEXT NOT NULL,
            fields TEXT NOT NULL,
            auto_width INTEGER NOT NULL DEFAULT 3,
            created_by TEXT,
            created_at TEXT NOT NULL
        )"""
    )
    conn.execute(
        """CREATE TABLE IF NOT EXISTS category_field_options (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            category_key TEXT NOT NULL,
            field_key TEXT NOT NULL,
            code TEXT NOT NULL,
            label TEXT NOT NULL,
            created_by TEXT,
            created_at TEXT NOT NULL,
            UNIQUE(category_key, field_key, code)
        )"""
    )
    existing = conn.execute("SELECT id FROM users WHERE username = ?", ("Master",)).fetchone()
    if not existing:
        conn.execute(
            "INSERT INTO users (username, password_hash, role, access, created_at) VALUES (?, ?, ?, ?, ?)",
            (
                "Master",
                generate_password_hash("Master@123"),
                "Admin",
                json.dumps(ALL_MODULES),
                datetime.utcnow().isoformat(timespec="seconds"),
            ),
        )
    conn.commit()
    conn.close()


def load_custom_categories(conn):
    """Admin-added material categories, stored as data (not code)."""
    out = {}
    for r in conn.execute("SELECT * FROM custom_categories").fetchall():
        fields_raw = json.loads(r["fields"])
        tuple_fields = [(f["key"], f["label"], f["kind"], f.get("choices")) for f in fields_raw]
        template = r["code_format"]

        def build_code(values, _template=template):
            return _template.format(**values)

        out[r["key"]] = {
            "key": r["key"],
            "label": r["label"],
            "code_format": template,
            "fields": tuple_fields,
            "scope_fields": None,  # sequence numbers run globally per category
            "auto_width": r["auto_width"],
            "build_code": build_code,
            "custom": True,
        }
    return out


def get_all_categories(conn):
    all_cats = dict(BUILTIN_CATEGORIES)
    all_cats.update(load_custom_categories(conn))
    return all_cats


# --------------------------------------------------------------------------
# Auth helpers
# --------------------------------------------------------------------------
def current_user():
    uid = session.get("uid")
    if not uid:
        return None
    conn = get_db()
    row = conn.execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
    conn.close()
    return row


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("uid"):
            if request.path.startswith("/api/"):
                return jsonify({"error": "not_authenticated"}), 401
            return redirect(url_for("login"))
        if current_user() is None:
            # Session points at a user id that no longer exists (DB was
            # reset, or the account was deleted while they were logged in).
            session.clear()
            if request.path.startswith("/api/"):
                return jsonify({"error": "not_authenticated"}), 401
            return redirect(url_for("login"))
        return f(*args, **kwargs)

    return wrapper


def admin_required(f):
    """Strictly Admin-role only — used for anything Admin status itself
    controls (creating material categories, granting Admin access)."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = current_user()
        if not user or user["role"] != "Admin":
            return jsonify({"error": "forbidden"}), 403
        return f(*args, **kwargs)

    return wrapper


def manage_users_required(f):
    """Admins always pass. A 'User' role account passes only if an Admin
    has explicitly granted them the 'user_management' module — but even
    then, granting Admin access itself stays gated by admin_required-style
    checks inside the route (see api_users / api_user_detail)."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = current_user()
        if not user:
            return jsonify({"error": "forbidden"}), 403
        access = json.loads(user["access"])
        if user["role"] != "Admin" and "user_management" not in access:
            return jsonify({"error": "forbidden"}), 403
        return f(*args, **kwargs)

    return wrapper


def user_to_dict(row):
    return {
        "id": row["id"],
        "username": row["username"],
        "role": row["role"],
        "access": json.loads(row["access"]),
        "created_at": row["created_at"],
    }


def slugify(text):
    slug = re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")
    return slug or "category"


# --------------------------------------------------------------------------
# Page routes
# --------------------------------------------------------------------------
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        if session.get("uid"):
            return redirect(url_for("index"))
        return render_template("login.html", company=COMPANY_NAME, error=None)

    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    conn = get_db()
    row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    conn.close()
    if not row or not check_password_hash(row["password_hash"], password):
        return render_template("login.html", company=COMPANY_NAME, error="Incorrect username or password.")
    session["uid"] = row["id"]
    return redirect(url_for("index"))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    user = current_user()
    return render_template(
        "app.html",
        company=COMPANY_NAME,
        plant=PLANT_INFO,
        user=user_to_dict(user),
        all_modules=ALL_MODULES,
        code_min=CODE_MIN_LEN,
        code_max=CODE_MAX_LEN,
    )


# --------------------------------------------------------------------------
# API: session / categories
# --------------------------------------------------------------------------
@app.route("/api/me")
@login_required
def api_me():
    return jsonify(user_to_dict(current_user()))


def _merged_choices(conn, category_key, field_key, base_choices):
    """Base (built-in or custom-category) choices, plus anything an Admin
    has added later via 'Manage options', flagged so the UI can tell which
    ones are removable."""
    result = [[c, l, False] for c, l in (base_choices or [])]
    existing = {c for c, l in (base_choices or [])}
    extra = conn.execute(
        "SELECT code, label FROM category_field_options WHERE category_key = ? AND field_key = ? ORDER BY id",
        (category_key, field_key),
    ).fetchall()
    for row in extra:
        if row["code"] in existing:
            continue
        result.append([row["code"], row["label"], True])
        existing.add(row["code"])
    return result


def category_length_range(cat, get_codes):
    """(shortest, longest) code length a category can ever produce, or None
    when a free-text field makes that unknowable. `get_codes(field_key,
    base_choices)` returns the list of option codes to consider for a
    dropdown field. Length is additive across segments, so the extremes
    come from picking the shortest / longest option in every field."""
    lo_vals, hi_vals = {}, {}
    for fkey, _flabel, fkind, fchoices in cat["fields"]:
        if fkind == "choice":
            codes = get_codes(fkey, fchoices)
            if not codes:
                return None
            lo_vals[fkey] = min(codes, key=len)
            hi_vals[fkey] = max(codes, key=len)
        elif fkind == "auto":
            lo_vals[fkey] = hi_vals[fkey] = "0" * cat.get("auto_width", 3)
        else:
            return None
    try:
        return len(cat["build_code"](lo_vals)), len(cat["build_code"](hi_vals))
    except KeyError:
        return None


def _length_rule_message(lo, hi):
    span = f"{lo}\u2013{hi} characters long" if lo != hi else f"{lo} characters long"
    return (f"Codes for this category would be {span} but every code must be "
            f"{CODE_MIN_LEN}\u2013{CODE_MAX_LEN} characters.")


def _categories_json(conn):
    out = {}
    for key, cat in get_all_categories(conn).items():
        fields_out = []
        for fkey, flabel, fkind, fchoices in cat["fields"]:
            if fkind == "choice":
                fchoices = _merged_choices(conn, key, fkey, fchoices)
            fields_out.append({"key": fkey, "label": flabel, "kind": fkind, "choices": fchoices})
        rng = category_length_range(
            cat, lambda fk, base, _k=key: [c[0] for c in _merged_choices(conn, _k, fk, base)]
        )
        out[key] = {
            "label": cat["label"],
            "code_format": cat["code_format"],
            "custom": cat.get("custom", False),
            "fields": fields_out,
            "length_range": list(rng) if rng else None,
        }
    # Imported historical codes may carry a category string that was never
    # formally defined (e.g. from a legacy spreadsheet). Surface it as a
    # read-only pseudo-category so History filters and labels still work.
    used = conn.execute("SELECT DISTINCT category FROM part_codes").fetchall()
    for row in used:
        key = row["category"]
        if key and key not in out:
            out[key] = {
                "label": key.replace("_", " ").replace("-", " ").title(),
                "code_format": "(imported / legacy)",
                "custom": True,
                "fields": [],
                "length_range": None,
            }
    return out


@app.route("/api/categories", methods=["GET", "POST"])
@login_required
def api_categories():
    conn = get_db()
    if request.method == "GET":
        out = _categories_json(conn)
        conn.close()
        return jsonify(out)

    # POST — create a new material category. Admin only: this is the one
    # thing that stays out of the delegable module-access system entirely.
    user = current_user()
    if user["role"] != "Admin":
        conn.close()
        return jsonify({"error": "Only an Admin can add a new material category."}), 403

    data = request.get_json(force=True)
    label = (data.get("label") or "").strip()
    prefix = re.sub(r"[^A-Za-z0-9-]", "", (data.get("prefix") or "").strip()) or "JEVBP"
    fields = data.get("fields") or []
    if not label:
        conn.close()
        return jsonify({"error": "Category name is required."}), 400
    if not fields:
        conn.close()
        return jsonify({"error": "Add at least one field."}), 400

    auto_fields = [f for f in fields if f.get("kind") == "auto"]
    if len(auto_fields) > 1:
        conn.close()
        return jsonify({"error": "Only one auto-sequence field is allowed per category."}), 400

    key = slugify(label)
    all_existing = get_all_categories(conn)
    suffix = 1
    base_key = key
    while key in all_existing:
        suffix += 1
        key = f"{base_key}_{suffix}"

    clean_fields = []
    template_parts = [prefix]
    for f in fields:
        fkey = slugify(f.get("key") or f.get("label") or "field")
        flabel = (f.get("label") or fkey).strip()
        fkind = f.get("kind") if f.get("kind") in ("choice", "text", "auto") else "text"
        choices = None
        if fkind == "choice":
            raw = f.get("choices_raw", "")
            choices = []
            for pair in raw.split(","):
                pair = pair.strip()
                if not pair:
                    continue
                if ":" in pair:
                    code, lbl = pair.split(":", 1)
                else:
                    code, lbl = pair, pair
                choices.append([code.strip(), lbl.strip()])
            if not choices:
                conn.close()
                return jsonify({"error": f"'{flabel}' needs at least one choice."}), 400
        clean_fields.append({"key": fkey, "label": flabel, "kind": fkind, "choices": choices})
        template_parts.append("{" + fkey + "}")

    code_format = "-".join(template_parts)
    auto_width = int(data.get("auto_width", 3))

    # Enforce the 10-18 character standard up front whenever it can be
    # worked out (i.e. no free-text field). Free-text categories are still
    # checked code-by-code at creation time.
    pseudo_cat = {
        "fields": [(f["key"], f["label"], f["kind"], f["choices"]) for f in clean_fields],
        "auto_width": auto_width,
        "build_code": lambda v, _t=code_format: _t.format(**v),
    }
    rng = category_length_range(pseudo_cat, lambda fk, base: [c[0] for c in (base or [])])
    if rng and (rng[0] < CODE_MIN_LEN or rng[1] > CODE_MAX_LEN):
        conn.close()
        return jsonify({"error": _length_rule_message(*rng) + " Shorten the prefix or option codes, or add more segments."}), 400

    conn.execute(
        "INSERT INTO custom_categories (key, label, code_format, fields, auto_width, created_by, created_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (key, label, code_format, json.dumps(clean_fields), auto_width,
         user["username"], datetime.utcnow().isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True, "key": key, "code_format": code_format})


@app.route("/api/categories/<key>", methods=["DELETE"])
@login_required
@admin_required
def api_delete_category(key):
    conn = get_db()
    row = conn.execute("SELECT key FROM custom_categories WHERE key = ?", (key,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Only Admin-added categories can be removed."}), 404
    in_use = conn.execute("SELECT id FROM part_codes WHERE category = ? LIMIT 1", (key,)).fetchone()
    if in_use:
        conn.close()
        return jsonify({"error": "This category already has codes issued against it and can't be removed."}), 400
    conn.execute("DELETE FROM custom_categories WHERE key = ?", (key,))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/categories/<key>/fields/<field_key>/options", methods=["POST"])
@login_required
@admin_required
def api_add_field_option(key, field_key):
    data = request.get_json(force=True)
    code = str(data.get("code", "")).strip()
    label = str(data.get("label", "")).strip() or code
    if not code:
        return jsonify({"error": "Option code is required."}), 400

    conn = get_db()
    categories = get_all_categories(conn)
    if key not in categories:
        conn.close()
        return jsonify({"error": "Unknown category."}), 404
    field = next((f for f in categories[key]["fields"] if f[0] == field_key), None)
    if not field or field[2] != "choice":
        conn.close()
        return jsonify({"error": "That field doesn't accept dropdown options."}), 400

    existing_codes = {c for c, l in (field[3] or [])}
    extra_rows = conn.execute(
        "SELECT code FROM category_field_options WHERE category_key = ? AND field_key = ?", (key, field_key)
    ).fetchall()
    existing_codes.update(r["code"] for r in extra_rows)
    if code in existing_codes:
        conn.close()
        return jsonify({"error": f"Option '{code}' already exists for this field."}), 409

    def _codes_with_new(fk, base, _key=key):
        codes = [c[0] for c in _merged_choices(conn, _key, fk, base)]
        if fk == field_key:
            codes.append(code)
        return codes

    rng = category_length_range(categories[key], _codes_with_new)
    if rng and (rng[0] < CODE_MIN_LEN or rng[1] > CODE_MAX_LEN):
        conn.close()
        return jsonify({"error": f"Can't add '{code}': " + _length_rule_message(*rng)}), 400

    conn.execute(
        "INSERT INTO category_field_options (category_key, field_key, code, label, created_by, created_at) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (key, field_key, code, label, current_user()["username"], datetime.utcnow().isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/categories/<key>/fields/<field_key>/options/<code>", methods=["DELETE"])
@login_required
@admin_required
def api_delete_field_option(key, field_key, code):
    conn = get_db()
    row = conn.execute(
        "SELECT id FROM category_field_options WHERE category_key = ? AND field_key = ? AND code = ?",
        (key, field_key, code),
    ).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Only options added here can be removed — this one is part of the base category definition."}), 400
    conn.execute("DELETE FROM category_field_options WHERE id = ?", (row["id"],))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


# --------------------------------------------------------------------------
# API: codes
# --------------------------------------------------------------------------
@app.route("/api/codes", methods=["GET", "POST"])
@login_required
def api_codes():
    conn = get_db()
    categories = get_all_categories(conn)

    if request.method == "POST":
        data = request.get_json(force=True)
        cat_key = data.get("category")
        if cat_key not in categories:
            conn.close()
            return jsonify({"error": "Unknown category."}), 400
        category = categories[cat_key]
        values = {}
        for fkey, flabel, fkind, _ in category["fields"]:
            if fkind == "auto":
                continue
            v = str(data.get("values", {}).get(fkey, "")).strip()
            if not v:
                conn.close()
                return jsonify({"error": f"'{flabel}' is required."}), 400
            values[fkey] = v

        for fkey, flabel, fkind, _ in category["fields"]:
            if fkind == "auto":
                scope_fields = category.get("scope_fields") or []
                scope_values = [values[f] for f in scope_fields]
                like_prefix = "-".join(scope_values) if scope_values else ""
                rows = conn.execute(
                    "SELECT code FROM part_codes WHERE category = ? ORDER BY id DESC", (cat_key,)
                ).fetchall()
                max_n = 0
                width = category.get("auto_width", 3)
                for r in rows:
                    code = r["code"]
                    if like_prefix and not code.startswith(like_prefix):
                        continue
                    tail = code.split("-")[-1]
                    if tail.isdigit():
                        max_n = max(max_n, int(tail))
                values[fkey] = str(max_n + 1).zfill(width)

        try:
            code = category["build_code"](values)
        except KeyError as e:
            conn.close()
            return jsonify({"error": f"Category template is missing a value for {e}."}), 400

        if not (CODE_MIN_LEN <= len(code) <= CODE_MAX_LEN):
            conn.close()
            return jsonify({
                "error": f"'{code}' is {len(code)} characters long; every code must be "
                         f"{CODE_MIN_LEN}\u2013{CODE_MAX_LEN} characters. Use shorter values."
            }), 400

        dup = conn.execute("SELECT id FROM part_codes WHERE code = ?", (code,)).fetchone()
        if dup:
            conn.close()
            return jsonify({
                "error": f"Duplicate blocked: '{code}' already exists in the register."
            }), 409

        user = current_user()
        created_at = datetime.utcnow().isoformat(timespec="seconds")
        cur = conn.execute(
            "INSERT INTO part_codes (category, code, attributes, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
            (cat_key, code, json.dumps(values), user["username"], created_at),
        )
        conn.commit()
        new_id = cur.lastrowid
        conn.close()
        return jsonify({
            "id": new_id, "code": code, "category": cat_key, "attributes": values,
            "created_by": user["username"], "created_at": created_at,
        })

    category_filter = request.args.get("category")
    q = "SELECT * FROM part_codes"
    params = []
    if category_filter and category_filter != "all":
        q += " WHERE category = ?"
        params.append(category_filter)
    q += " ORDER BY id DESC"
    rows = conn.execute(q, params).fetchall()
    conn.close()
    return jsonify([
        {
            "id": r["id"], "category": r["category"], "code": r["code"],
            "attributes": json.loads(r["attributes"]), "created_by": r["created_by"],
            "created_at": r["created_at"],
        }
        for r in rows
    ])


@app.route("/api/stats")
@login_required
def api_stats():
    granularity = request.args.get("granularity", "day")
    fmt = {"day": "%Y-%m-%d", "month": "%Y-%m", "year": "%Y"}.get(granularity, "%Y-%m-%d")
    conn = get_db()
    rows = conn.execute("SELECT created_at FROM part_codes").fetchall()
    conn.close()
    buckets = {}
    for r in rows:
        try:
            dt = datetime.fromisoformat(r["created_at"])
        except ValueError:
            continue
        key = dt.strftime(fmt)
        buckets[key] = buckets.get(key, 0) + 1
    ordered = sorted(buckets.items())
    return jsonify({"labels": [k for k, _ in ordered], "values": [v for _, v in ordered]})


# --------------------------------------------------------------------------
# API: export (txt / xlsx)
# --------------------------------------------------------------------------
def _xlsx_col(idx):
    """0-based column index -> spreadsheet letters (0 -> A, 26 -> AA)."""
    idx += 1
    letters = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def build_xlsx(headers, rows, sheet_name="Part Codes"):
    """Return the bytes of a real .xlsx workbook using only the standard
    library, so Excel export works with nothing but Flask installed."""
    def cell(ref, value, style=0):
        text = "" if value is None else str(value)
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        s = f' s="{style}"' if style else ""
        if isinstance(value, int) and not isinstance(value, bool):
            return f'<c r="{ref}"{s}><v>{value}</v></c>'
        return f'<c r="{ref}" t="inlineStr"{s}><is><t xml:space="preserve">{xml_escape(text)}</t></is></c>'

    everything = [headers] + rows
    widths = [
        max(len(str(r[i])) if i < len(r) and r[i] is not None else 0 for r in everything)
        for i in range(len(headers))
    ]
    cols = "".join(
        f'<col min="{i + 1}" max="{i + 1}" width="{min(max(w + 2, 10), 50)}" customWidth="1"/>'
        for i, w in enumerate(widths)
    )
    sheet_rows = []
    for r_idx, row in enumerate(everything, start=1):
        style = 1 if r_idx == 1 else 0
        cells = "".join(cell(f"{_xlsx_col(c)}{r_idx}", v, style) for c, v in enumerate(row))
        sheet_rows.append(f'<row r="{r_idx}">{cells}</row>')

    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_rel = "http://schemas.openxmlformats.org/package/2006/relationships"
    head = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'

    sheet_xml = (
        f'{head}<worksheet xmlns="{ns}"><sheetViews><sheetView workbookViewId="0">'
        '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        f'</sheetView></sheetViews><cols>{cols}</cols><sheetData>{"".join(sheet_rows)}</sheetData></worksheet>'
    )
    styles_xml = (
        f'{head}<styleSheet xmlns="{ns}">'
        '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
        '<font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/></font></fonts>'
        '<fills count="3"><fill><patternFill patternType="none"/></fill>'
        '<fill><patternFill patternType="gray125"/></fill>'
        '<fill><patternFill patternType="solid"><fgColor rgb="FF034078"/><bgColor indexed="64"/></patternFill></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/></cellXfs>'
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>'
    )
    workbook_xml = (
        f'{head}<workbook xmlns="{ns}" xmlns:r="{rel_ns}"><sheets>'
        f'<sheet name="{xml_escape(sheet_name)}" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    workbook_rels = (
        f'{head}<Relationships xmlns="{pkg_rel}">'
        f'<Relationship Id="rId1" Type="{rel_ns}/worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{rel_ns}/styles" Target="styles.xml"/></Relationships>'
    )
    root_rels = (
        f'{head}<Relationships xmlns="{pkg_rel}">'
        f'<Relationship Id="rId1" Type="{rel_ns}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    content_types = (
        f'{head}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '</Types>'
    )

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("xl/workbook.xml", workbook_xml)
        z.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        z.writestr("xl/styles.xml", styles_xml)
        z.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return buf.getvalue()


@app.route("/api/export", methods=["POST"])
@login_required
def api_export():
    data = request.get_json(force=True)
    ids = data.get("ids") or []
    fmt = data.get("format", "txt")
    if not ids:
        return jsonify({"error": "No codes selected."}), 400

    conn = get_db()
    placeholders = ",".join("?" for _ in ids)
    rows = conn.execute(
        f"SELECT * FROM part_codes WHERE id IN ({placeholders}) ORDER BY id DESC", ids
    ).fetchall()
    categories = get_all_categories(conn)
    conn.close()

    if fmt == "xlsx":
        table = []
        for r in rows:
            cat_label = categories.get(r["category"], {}).get("label", r["category"])
            attrs = json.loads(r["attributes"])
            attrs_str = "; ".join(f"{k}={v}" for k, v in attrs.items())
            table.append([r["code"], len(r["code"]), cat_label, r["created_by"], r["created_at"], attrs_str])
        buf = io.BytesIO(build_xlsx(
            ["Code", "Length", "Category", "Created By", "Created At (UTC)", "Attributes"], table
        ))
        return send_file(
            buf, as_attachment=True, download_name="part_codes_export.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # default: txt
    content = "\n".join(r["code"] for r in rows)
    buf = io.BytesIO(content.encode("utf-8"))
    return send_file(buf, as_attachment=True, download_name="part_codes_export.txt", mimetype="text/plain")


@app.route("/api/import", methods=["POST"])
@login_required
@admin_required
def api_import():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Choose a .csv or .xlsx file first."}), 400

    filename = file.filename
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    raw_rows = []

    if ext == "csv":
        content = file.stream.read().decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(content))
        for row in reader:
            raw_rows.append(row)
    elif ext in ("xlsx", "xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            return jsonify({
                "error": "Reading .xlsx files needs the 'openpyxl' package (pip install openpyxl). "
                         "Or save the sheet as .csv and import that instead."
            }), 500
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        headers = None
        for row in ws.iter_rows(values_only=True):
            if headers is None:
                headers = [str(h).strip() if h is not None else "" for h in row]
                continue
            raw_rows.append({headers[i]: row[i] for i in range(len(headers)) if i < len(row)})
    else:
        return jsonify({"error": "Unsupported file type — upload a .csv or .xlsx file."}), 400

    # Normalize header keys (case/whitespace-insensitive lookup)
    norm_rows = []
    for r in raw_rows:
        norm_rows.append({(k.strip().lower() if k else k): v for k, v in r.items()})

    conn = get_db()
    inserted = skipped_dup = skipped_invalid = off_standard = 0

    for r in norm_rows:
        code = str(r.get("code") or "").strip()
        if not code:
            skipped_invalid += 1
            continue
        category = str(r.get("category") or "legacy").strip() or "legacy"
        created_by = str(r.get("created_by") or "import").strip() or "import"
        created_at_raw = str(r.get("created_at") or "").strip()
        try:
            created_at = (
                datetime.fromisoformat(created_at_raw).isoformat(timespec="seconds")
                if created_at_raw else datetime.utcnow().isoformat(timespec="seconds")
            )
        except ValueError:
            created_at = datetime.utcnow().isoformat(timespec="seconds")

        dup = conn.execute("SELECT id FROM part_codes WHERE code = ?", (code,)).fetchone()
        if dup:
            skipped_dup += 1
            continue

        conn.execute(
            "INSERT INTO part_codes (category, code, attributes, created_by, created_at) VALUES (?, ?, ?, ?, ?)",
            (category, code, json.dumps({"imported": True}), created_by, created_at),
        )
        inserted += 1
        if not (CODE_MIN_LEN <= len(code) <= CODE_MAX_LEN):
            off_standard += 1

    conn.commit()
    conn.close()
    return jsonify({
        "ok": True, "inserted": inserted, "skipped_duplicate": skipped_dup,
        "skipped_invalid": skipped_invalid, "total_rows": len(norm_rows),
        "off_standard": off_standard,
    })


# --------------------------------------------------------------------------
# API: user management
# --------------------------------------------------------------------------
@app.route("/api/users", methods=["GET", "POST"])
@login_required
@manage_users_required
def api_users():
    conn = get_db()
    acting_user = current_user()

    if request.method == "POST":
        data = request.get_json(force=True)
        username = data.get("username", "").strip()
        password = data.get("password", "")
        requested_role = data.get("role", "User")
        access = data.get("access", [])

        if not username or not password:
            conn.close()
            return jsonify({"error": "Username and password are required."}), 400

        # Only a real Admin can create another Admin — non-admin managers
        # (delegated user_management access) are forced to "User".
        if requested_role == "Admin" and acting_user["role"] != "Admin":
            conn.close()
            return jsonify({"error": "Only an Admin can grant Admin access."}), 403
        role = "Admin" if requested_role == "Admin" else "User"

        access = [a for a in access if a in ALL_MODULES]
        if role == "Admin":
            access = ALL_MODULES[:]
        elif acting_user["role"] != "Admin":
            access = [a for a in access if a != "user_management"]

        existing = conn.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone()
        if existing:
            conn.close()
            return jsonify({"error": "Username already exists."}), 409

        conn.execute(
            "INSERT INTO users (username, password_hash, role, access, created_at) VALUES (?, ?, ?, ?, ?)",
            (username, generate_password_hash(password), role, json.dumps(access),
             datetime.utcnow().isoformat(timespec="seconds")),
        )
        conn.commit()
        conn.close()
        return jsonify({"ok": True})

    rows = conn.execute("SELECT * FROM users ORDER BY id ASC").fetchall()
    conn.close()
    return jsonify([user_to_dict(r) for r in rows])


@app.route("/api/users/<int:uid>", methods=["PUT", "DELETE"])
@login_required
@manage_users_required
def api_user_detail(uid):
    conn = get_db()
    acting_user = current_user()
    target = conn.execute("SELECT * FROM users WHERE id = ?", (uid,)).fetchone()
    if not target:
        conn.close()
        return jsonify({"error": "User not found."}), 404

    if request.method == "DELETE":
        if target["username"] == "Master":
            conn.close()
            return jsonify({"error": "The Master admin account can't be deleted."}), 400
        conn.execute("DELETE FROM users WHERE id = ?", (uid,))
        conn.commit()
        conn.close()
        return jsonify({"ok": True})

    data = request.get_json(force=True)
    requested_role = data.get("role", target["role"])
    if requested_role == "Admin" and acting_user["role"] != "Admin":
        conn.close()
        return jsonify({"error": "Only an Admin can grant Admin access."}), 403
    role = "Admin" if requested_role == "Admin" else "User"

    access = data.get("access", json.loads(target["access"]))
    access = [a for a in access if a in ALL_MODULES]
    if role == "Admin":
        access = ALL_MODULES[:]
    elif acting_user["role"] != "Admin":
        access = [a for a in access if a != "user_management"]

    conn.execute("UPDATE users SET role = ?, access = ? WHERE id = ?", (role, json.dumps(access), uid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


@app.route("/api/users/<int:uid>/reset_password", methods=["POST"])
@login_required
@manage_users_required
def api_reset_password(uid):
    data = request.get_json(force=True)
    new_password = data.get("password", "")
    if not new_password or len(new_password) < 4:
        return jsonify({"error": "New password must be at least 4 characters."}), 400
    conn = get_db()
    conn.execute("UPDATE users SET password_hash = ? WHERE id = ?", (generate_password_hash(new_password), uid))
    conn.commit()
    conn.close()
    return jsonify({"ok": True})


if __name__ == "__main__":
    init_db()
    app.run(debug=True, host="0.0.0.0", port=5000)
